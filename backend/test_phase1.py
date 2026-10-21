import os
import json
import pytest
import docx
from pathlib import Path
from extractor import extract_protocol

ARTIFACTS_DIR = Path(__file__).resolve().parent / "artifacts"
UPLOADS_DIR = Path(__file__).resolve().parent / "uploads"

def test_dom_map_verification():
    dom_file = ARTIFACTS_DIR / "dom_map.json"
    assert dom_file.exists(), "dom_map.json does not exist!"
    
    with open(dom_file, "r", encoding="utf-8") as f:
        data = json.load(f)
        
    assert data["target_url"] == "https://mpv-two.vercel.app"
    assert "login_page" in data
    assert "header_shell" in data
    assert "manual_case_registration" in data
    
    # Check verified login selectors
    login_keys = data["login_page"]["key_selectors"]
    assert login_keys["username_input"] == "#username"
    assert login_keys["password_input"] == "#password"
    assert login_keys["enterprise_input"] == "#enterprise"
    assert login_keys["login_button"] == "button[type='submit']"
    
    # Check verified header elements
    shell = data["header_shell"]
    assert shell["header_username_found"] is True
    assert shell["header_timestamp_found"] is True
    assert shell["case_exit_found"] is True
    assert shell["simulation_bar_found"] is True
    print("[PASS] dom_map.json verified successfully with live target selectors.")

def test_pq022_extraction():
    pq_path = UPLOADS_DIR / "CLIN-EU-PAT-OPV-1-1-PQ-022- Manual - Registering Adverse Event Case.docx"
    assert pq_path.exists(), f"PQ-022 file missing at {pq_path}"
    
    session = extract_protocol(str(pq_path))
    assert session.document_no == "CLIN-EU-PAT-OPV-1-1-PQ-022"
    assert session.total_steps == 18, f"Expected 18 steps, got {session.total_steps}"
    assert session.total_screenshots == 18, f"Expected 18 screenshots, got {session.total_screenshots}"
    assert session.steps[0].step_id == "TS1-S01"
    assert session.steps[-1].step_id == "TS1-S18"
    assert session.steps[0].action_intent == "NAVIGATE_LOGIN"
    assert session.steps[1].action_intent == "LOGIN_ONEPV"
    assert session.steps[2].action_intent == "NAVIGATE_MANUAL_CASE"
    print(f"[PASS] PQ-022 extracted: {session.total_steps} steps, {session.total_screenshots} screenshots.")

def test_oq036_extraction():
    oq_path = UPLOADS_DIR / "CLIN-EU-VAL-OPV-1-1-OQ-036 - Manual - Registering Adverse Event Case.docx"
    assert oq_path.exists(), f"OQ-036 file missing at {oq_path}"
    
    session = extract_protocol(str(oq_path))
    assert session.document_no == "CLIN-EU-VAL-OPV-1-1-OQ-036"
    assert session.total_steps == 50, f"Expected 50 steps, got {session.total_steps}"
    assert session.total_screenshots == 50, f"Expected 50 screenshots, got {session.total_screenshots}"
    assert session.steps[0].step_id == "TS1-S01"
    assert session.steps[-1].step_id == "TS2-S08"
    print(f"[PASS] OQ-036 extracted: {session.total_steps} steps, {session.total_screenshots} screenshots.")

def test_synthetic_protocol_with_footers_and_merges(tmp_path):
    # Create synthetic docx
    doc = docx.Document()
    doc.add_paragraph("CLIN-SYNTH-999 Validation Protocol")
    doc.add_paragraph("1.0 OBJECTIVE")
    doc.add_paragraph("Synthetic test objective for automated verification.")
    doc.add_paragraph("2.0 SCOPE")
    doc.add_paragraph("Synthetic test scope covering adverse event workflows.")
    
    table = doc.add_table(rows=5, cols=5)
    # Row 0: Header
    table.rows[0].cells[0].text = "Step #"
    table.rows[0].cells[1].text = "Scenario"
    table.rows[0].cells[2].text = "Procedure"
    table.rows[0].cells[3].text = "Expected Result"
    table.rows[0].cells[4].text = "Initials / Date"
    
    # Row 1: Step with micro symbol and multi assertion
    table.rows[1].cells[0].text = ""  # Test auto-numbering
    table.rows[1].cells[1].text = "Register Adverse Event Case"
    table.rows[1].cells[2].text = "Fill form and click Save. (Refer expected result 1) \u00b5 Click on OK button. (Refer expected result 2)"
    table.rows[1].cells[3].text = "Case ID created and dialog displayed."
    
    # Row 2: Step with inherited scenario (empty scenario cell)
    table.rows[2].cells[0].text = ""
    table.rows[2].cells[1].text = ""  # Should inherit from row 1
    table.rows[2].cells[2].text = "Submit the created case for review. \u00b5"
    table.rows[2].cells[3].text = "Case submitted successfully."
    
    # Row 3: Merged Footer Row Comments / Deviations
    table.rows[3].cells[0].text = "Comments / Deviations: None observed during test."
    table.rows[3].cells[1].text = "Comments / Deviations: None observed during test."
    
    # Row 4: Footer Row Executed By
    table.rows[4].cells[0].text = "Executed By: S.H. / 14-Oct-2026"
    
    synth_path = tmp_path / "CLIN-SYNTH-999.docx"
    doc.save(str(synth_path))
    
    session = extract_protocol(str(synth_path))
    assert session.document_no == "CLIN-SYNTH-999"
    assert session.total_steps == 2, f"Footer rows should be excluded! Got {session.total_steps} steps"
    assert session.steps[0].step_id == "TS1-S01"
    assert session.steps[1].step_id == "TS1-S02"
    assert session.steps[0].screenshot_required is True
    assert session.steps[0].capture_before_modal_ok is True
    assert session.steps[0].required_assertions_count == 2
    assert session.steps[1].scenario == "Register Adverse Event Case", "Should have inherited scenario"
    print("[PASS] Synthetic protocol with footers and merged cells verified.")

if __name__ == "__main__":
    pytest.main(["-v", __file__])
